"""Gera reports/parcial/relatorio_parcial_primeira_entrega.docx.

Script de geração reprodutível do relatório parcial (Primeira Entrega, cronograma
oficial — vencimento 01/09/2026, carência 06/09/2026), seguindo a estrutura do
template institucional em modelos_anteriores_tcc/Modelo_Projeto_TCC_Científico.docx
e as convenções de formatação observadas nos exemplos de TCCs anteriores do curso de
Ciência de Dados (capa, folha de rosto, resumo, sumário, fundamentação teórica,
metodologia, cronograma, referências).

Não reutiliza tema, dados ou autoria dos exemplos anteriores — apenas a estrutura e
as regras de formatação (Times New Roman, espaçamento 1,5 no corpo, títulos de
capítulo em maiúsculas/negrito iniciando nova página, ABNT NBR 14724/6023).

Uso: `python3 reports/parcial/gerar_relatorio_parcial.py`
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent.parent
SAIDA = Path(__file__).resolve().parent / "relatorio_parcial_primeira_entrega.docx"

TITULO_TRABALHO = (
    "Apostas online e vulnerabilidade financeira das famílias brasileiras: "
    "mensuração nacional, quebras estruturais e projeção de cenários (2020–2029)"
)

# Ordem alfabética por primeiro nome, mesma convenção observada nos exemplos do curso.
AUTORES = [
    "Fabiano Guilherme Dionizio Bortolussi",
    "José Araújo Neto",
    "Josué Dantas Martins",
    "Palmira da Conceição João",
    "Pedrina Ferreira Gomes",
    "Vinícius Figueiredo Dias Nunes",
    "Wilton da Silva Alves",
    "Yan Ferreira Martins",
]
ORIENTADOR = "Felipe Ivo da Silva"
AVALIADOR_EXTERNO = "Eduardo Noriyuki Sakuma Shibata"
CIDADE_ANO = "São Paulo – SP\n2026"


# --------------------------------------------------------------------------- #
# Configuração de estilos base (ABNT NBR 14724, conforme template institucional)
# --------------------------------------------------------------------------- #

def configurar_documento(doc: Document) -> None:
    estilo_normal = doc.styles["Normal"]
    estilo_normal.font.name = "Times New Roman"
    estilo_normal.font.size = Pt(12)
    rpr = estilo_normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "Times New Roman")

    fmt = estilo_normal.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_after = Pt(0)

    for secao in doc.sections:
        secao.top_margin = Cm(3)
        secao.left_margin = Cm(3)
        secao.bottom_margin = Cm(2)
        secao.right_margin = Cm(2)


def paragrafo(doc, texto="", *, tamanho=12, negrito=False, italico=False,
              alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, espacamento_simples=False,
              recuo_primeira_linha=None, recuo_esquerdo=None):
    p = doc.add_paragraph()
    p.alignment = alinhamento
    if espacamento_simples:
        p.paragraph_format.line_spacing = 1.0
    if recuo_primeira_linha is not None:
        p.paragraph_format.first_line_indent = Cm(recuo_primeira_linha)
    if recuo_esquerdo is not None:
        p.paragraph_format.left_indent = Cm(recuo_esquerdo)
    if texto:
        run = p.add_run(texto)
        run.font.size = Pt(tamanho)
        run.font.bold = negrito
        run.font.italic = italico
    return p


def titulo_capitulo(doc, texto: str, nova_pagina: bool = True):
    if nova_pagina:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(texto.upper())
    run.font.bold = True
    run.font.size = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    return p


def subtitulo(doc, texto: str, nivel: int = 1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(texto.upper() if nivel == 1 else texto)
    run.font.bold = True
    run.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p


def tabela_simples(doc, cabecalho: list[str], linhas: list[list[str]]):
    tabela = doc.add_table(rows=1, cols=len(cabecalho))
    tabela.style = "Light Grid Accent 1"
    for i, texto in enumerate(cabecalho):
        celula = tabela.rows[0].cells[i]
        celula.text = ""
        run = celula.paragraphs[0].add_run(texto)
        run.font.bold = True
        run.font.size = Pt(10)
    for linha in linhas:
        cels = tabela.add_row().cells
        for i, texto in enumerate(linha):
            cels[i].text = ""
            run = cels[i].paragraphs[0].add_run(texto)
            run.font.size = Pt(10)
    doc.add_paragraph()
    return tabela


def referencia(doc, texto: str):
    paragrafo(doc, texto, tamanho=12, alinhamento=WD_ALIGN_PARAGRAPH.LEFT, espacamento_simples=True)
    p = doc.paragraphs[-1]
    p.paragraph_format.space_after = Pt(12)


# --------------------------------------------------------------------------- #
# Construção do documento
# --------------------------------------------------------------------------- #

def construir_capa(doc):
    for _ in range(3):
        doc.add_paragraph()
    paragrafo(doc, "UNIVERSIDADE VIRTUAL DO ESTADO DE SÃO PAULO", negrito=True,
              alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    paragrafo(doc, "BACHARELADO EM CIÊNCIA DE DADOS", negrito=True,
              alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(3):
        doc.add_paragraph()
    for autor in AUTORES:
        paragrafo(doc, autor, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(4):
        doc.add_paragraph()
    paragrafo(doc, TITULO_TRABALHO, negrito=True, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    paragrafo(doc, "Relatório Parcial — Primeira Entrega", italico=True,
              alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(6):
        doc.add_paragraph()
    for linha in CIDADE_ANO.split("\n"):
        paragrafo(doc, linha, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)


def construir_folha_de_rosto(doc):
    doc.add_page_break()
    for _ in range(2):
        doc.add_paragraph()
    for autor in AUTORES:
        paragrafo(doc, autor, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(3):
        doc.add_paragraph()
    paragrafo(doc, TITULO_TRABALHO, negrito=True, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    paragrafo(doc, "Relatório Parcial — Primeira Entrega", italico=True,
              alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(3):
        doc.add_paragraph()
    paragrafo(
        doc,
        "Relatório parcial apresentado como Primeira Entrega do Trabalho de Conclusão "
        "de Curso do Bacharelado em Ciência de Dados da Universidade Virtual do "
        "Estado de São Paulo (UNIVESP), referente à Quinzena 2 do cronograma oficial "
        "de Trabalhos de Conclusão de Curso, com vencimento em 01/09/2026 e carência "
        "até 06/09/2026.",
        alinhamento=WD_ALIGN_PARAGRAPH.RIGHT, recuo_esquerdo=8,
    )
    for _ in range(2):
        doc.add_paragraph()
    paragrafo(doc, f"Orientador: {ORIENTADOR}", alinhamento=WD_ALIGN_PARAGRAPH.RIGHT, recuo_esquerdo=8)
    paragrafo(doc, f"Avaliador Externo: {AVALIADOR_EXTERNO}", alinhamento=WD_ALIGN_PARAGRAPH.RIGHT, recuo_esquerdo=8)
    for _ in range(6):
        doc.add_paragraph()
    for linha in CIDADE_ANO.split("\n"):
        paragrafo(doc, linha, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)


def construir_ficha(doc):
    doc.add_page_break()
    for _ in range(8):
        doc.add_paragraph()
    entrada = (
        "BORTOLUSSI, Fabiano Guilherme Dionizio; NETO, José Araújo; MARTINS, Josué "
        "Dantas; JOÃO, Palmira da Conceição; GOMES, Pedrina Ferreira; NUNES, "
        "Vinícius Figueiredo Dias; ALVES, Wilton da Silva; MARTINS, Yan Ferreira. "
        f"{TITULO_TRABALHO}. Relatório Parcial (Primeira Entrega) de Trabalho de "
        "Conclusão de Curso — Bacharelado em Ciência de Dados, Universidade Virtual "
        f"do Estado de São Paulo. Orientador: {ORIENTADOR}. São Paulo, 2026."
    )
    paragrafo(doc, entrada, tamanho=11, recuo_esquerdo=8)


def construir_resumo(doc):
    titulo_capitulo(doc, "RESUMO")
    texto = (
        "Este relatório apresenta o progresso do Trabalho de Conclusão de Curso do "
        "Bacharelado em Ciência de Dados da Universidade Virtual do Estado de São "
        "Paulo (UNIVESP) referente à Primeira Entrega do cronograma oficial "
        "(Quinzena 2). O trabalho investiga a relação entre a expansão das apostas "
        "online e a vulnerabilidade financeira das famílias brasileiras, com foco na "
        "mensuração nacional de indicadores de exposição a apostas e de endividamento/"
        "inadimplência, na detecção de quebras estruturais nas séries financeiras "
        "nacionais associadas aos marcos regulatórios de 2025, e na projeção de "
        "cenários de endividamento em horizonte de 24 meses. A pergunta central de "
        "pesquisa é: qual a trajetória projetada dos indicadores de endividamento e "
        "de inadimplência das famílias brasileiras sob a manutenção do padrão atual "
        "de exposição a apostas online? A unidade de análise principal é o Brasil, "
        "com aprofundamentos exploratórios por Unidade da Federação e por município. "
        "A metodologia integra fontes públicas do Banco Central do Brasil (SGS, "
        "EPAE), da Confederação Nacional do Comércio (PEIC), da Secretaria de "
        "Prêmios e Apostas do Ministério da Fazenda, do IBGE (via API SIDRA), do "
        "Portal da Transparência e do Google Trends, tratadas sob os princípios de "
        "reprodutibilidade, ausência de vazamento de dados entre treino e teste "
        "(validação exclusivamente por origem móvel) e transparência sobre "
        "limitações metodológicas. Até o momento desta entrega, foram concluídos: a "
        "especificação completa do projeto por meio de desenvolvimento orientado por "
        "especificações, a definição da arquitetura técnica, e a implementação e "
        "validação, por meio de 161 testes automatizados, de 27 módulos de software "
        "cobrindo ingestão de dados (com três fontes já integradas a interfaces de "
        "programação de aplicação oficiais — Banco Central, IBGE/SIDRA e Portal da "
        "Transparência), tratamento estatístico, diagnóstico exploratório, detecção "
        "de quebras estruturais, construção de indicador composto territorial, "
        "modelagem preditiva comparativa, avaliação estatística de modelos e geração "
        "de relatórios e painel interativo. As próximas etapas, a serem "
        "desenvolvidas até a Segunda Entrega, envolvem a execução do pipeline sobre "
        "dados reais coletados e o aprofundamento da revisão de literatura."
    )
    paragrafo(doc, texto, recuo_primeira_linha=1.25)
    doc.add_paragraph()
    p = paragrafo(doc, "", alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run("Palavras-chave: ")
    run.font.bold = True
    p.add_run(
        "Apostas online; Endividamento das famílias; Séries temporais; Quebras "
        "estruturais; Ciência de Dados."
    )


def construir_sumario(doc):
    titulo_capitulo(doc, "SUMÁRIO")
    itens = [
        "1 INTRODUÇÃO",
        "1.1 Contexto",
        "1.2 Problema de Pesquisa",
        "1.3 Objetivos",
        "1.4 Justificativa",
        "1.5 Estrutura deste Relatório",
        "2 FUNDAMENTAÇÃO TEÓRICA PRELIMINAR",
        "2.1 Apostas Online e Vulnerabilidade Financeira: Evidências Internacionais",
        "2.2 Apostas Online e Endividamento das Famílias Brasileiras: Evidências Nacionais",
        "2.3 O Marco Regulatório Brasileiro",
        "2.4 Detecção de Quebras Estruturais em Séries Temporais Econômicas",
        "3 METODOLOGIA",
        "3.1 Tipo de Pesquisa",
        "3.2 Recorte e Unidade de Análise",
        "3.3 Fontes de Dados e Estratégia de Coleta",
        "3.4 Tratamento e Análise Exploratória",
        "3.5 Detecção de Quebras Estruturais e Validação Externa",
        "3.6 Modelagem Preditiva e Validação",
        "3.7 Projeção de Cenários e Custo de Oportunidade Patrimonial",
        "3.8 Ferramentas e Tecnologias",
        "3.9 Estado Atual da Implementação",
        "3.10 Princípios Metodológicos Inegociáveis",
        "4 CRONOGRAMA",
        "REFERÊNCIAS",
    ]
    for item in itens:
        nivel_2 = "." in item.split(" ")[0]
        p = paragrafo(doc, item, alinhamento=WD_ALIGN_PARAGRAPH.LEFT, espacamento_simples=True,
                       recuo_esquerdo=1.25 if nivel_2 else None)
        p.paragraph_format.space_after = Pt(4)


# --- Capítulo 1 --------------------------------------------------------- #

def construir_introducao(doc):
    titulo_capitulo(doc, "1 INTRODUÇÃO")

    subtitulo(doc, "1.1 Contexto")
    paragrafo(doc, (
        "O mercado de apostas online tornou-se, em poucos anos, um dos maiores "
        "canais de saída de renda das famílias brasileiras, e sua expansão "
        "coincidiu com um período de endividamento historicamente elevado. Estudo "
        "técnico do Banco Central do Brasil identificou que, apenas em agosto de "
        "2024, pessoas pertencentes a famílias beneficiárias do Programa Bolsa "
        "Família transferiram cerca de R$ 3 bilhões a empresas de apostas por meio "
        "do Pix, com valor mediano de R$ 100 por indivíduo, sendo que aproximadamente "
        "17% dos cadastrados apostaram no período (BANCO CENTRAL DO BRASIL, 2024). A "
        "Secretaria de Prêmios e Apostas do Ministério da Fazenda registrou 17,7 "
        "milhões de apostadores no primeiro semestre de 2025, receita bruta de jogo "
        "de R$ 17,4 bilhões e gasto médio de aproximadamente R$ 164 mensais por "
        "apostador ativo (SPA/MF, 2025). O Boletim Fiscal dos Estados Brasileiros "
        "estimou perda líquida de R$ 62,5 bilhões das famílias brasileiras para "
        "plataformas de apostas, associada a um fluxo de aproximadamente R$ 351 "
        "bilhões em transferências via Pix (COMSEFAZ; CENTRO INTERNACIONAL CELSO "
        "FURTADO, 2026)."
    ), recuo_primeira_linha=1.25)
    paragrafo(doc, (
        "O tema é cientificamente tratável, e não apenas socialmente relevante, "
        "pela existência de três marcos regulatórios com data definida que "
        "segmentam a série histórica e permitem observar mudanças de regime no "
        "comportamento agregado: a entrada em vigor da exigência de autorização das "
        "operadoras em janeiro de 2025 (Lei n.º 14.790/2023); a vedação da "
        "participação de beneficiários do Bolsa Família e do BPC em outubro de 2025 "
        "(Portaria SPA/MF n.º 2.217/2025 e Instrução Normativa SPA/MF n.º 22/2025); "
        "e a suspensão parcial, por decisão do Supremo Tribunal Federal, das "
        "obrigações de bloqueio e encerramento de contas em dezembro de 2025, com "
        "manutenção da vedação a novos cadastros."
    ), recuo_primeira_linha=1.25)

    subtitulo(doc, "1.2 Problema de Pesquisa")
    paragrafo(doc, (
        "A lacuna que este trabalho pretende ocupar é de natureza empírica: os "
        "estudos disponíveis são pontuais, publicados de forma dispersa e não "
        "reprodutíveis. Não existe base pública consolidada que reúna, em série "
        "histórica única e documentada, indicadores de exposição a apostas e de "
        "vulnerabilidade financeira das famílias brasileiras. A pergunta central "
        "que orienta a pesquisa é: qual a trajetória projetada dos indicadores de "
        "endividamento e de inadimplência das famílias brasileiras sob a "
        "manutenção do padrão atual de exposição a apostas online?"
    ), recuo_primeira_linha=1.25)
    paragrafo(doc, "Dela decorrem quatro perguntas operacionais:", recuo_primeira_linha=1.25)
    for texto in [
        "Mensuração — é possível construir, exclusivamente a partir de fontes públicas, um indicador nacional de exposição a apostas online consistente e reprodutível?",
        "Detecção — as séries financeiras nacionais registram quebras estruturais e, em caso positivo, essas quebras coincidem com os marcos regulatórios de 2025?",
        "Projeção — mantido o padrão atual de comportamento, qual a trajetória esperada do endividamento das famílias no horizonte de 24 meses, e que magnitude assume o custo de oportunidade patrimonial em horizontes mais longos?",
        "Distribuição territorial — que Unidade da Federação apresenta maior exposição relativa segundo o indicador composto proposto, e como se distribui essa exposição entre seus principais municípios?",
    ]:
        paragrafo(doc, f"• {texto}", recuo_esquerdo=1.25)

    subtitulo(doc, "1.3 Objetivos")
    paragrafo(doc, (
        "Objetivo geral: construir uma base de dados integrada e um conjunto de "
        "modelos preditivos que descrevam a evolução da exposição a apostas online "
        "e da vulnerabilidade financeira das famílias brasileiras, projetando "
        "cenários para o período subsequente."
    ), recuo_primeira_linha=1.25)
    paragrafo(doc, "Objetivos específicos:", recuo_primeira_linha=1.25)
    for texto in [
        "Integrar fontes públicas heterogêneas em uma base nacional mensal reprodutível;",
        "Construir e validar um indicador de exposição a apostas online;",
        "Detectar quebras estruturais nas séries e confrontá-las com o calendário regulatório;",
        "Comparar modelos de previsão de séries temporais sob validação temporal rigorosa;",
        "Projetar cenários e estimar, por simulação, o custo de oportunidade patrimonial associado;",
        "Identificar, por indicador composto, a Unidade da Federação de maior exposição relativa e ilustrar sua distribuição municipal interna;",
        "Publicar base, código e documentação sob princípios de dados abertos.",
    ]:
        paragrafo(doc, f"• {texto}", recuo_esquerdo=1.25)

    subtitulo(doc, "1.4 Justificativa")
    paragrafo(doc, (
        "A relevância social do recorte é direta: as evidências disponíveis "
        "indicam participação desproporcional de famílias de menor renda no "
        "mercado de apostas online. Trata-se de um vetor de endividamento que "
        "incide justamente sobre o estrato populacional com menor capacidade de "
        "absorção de perdas e menor margem para acumulação patrimonial. Do ponto "
        "de vista acadêmico, o trabalho contribui ao integrar, em uma base única, "
        "documentada e reprodutível, indicadores hoje dispersos entre boletins "
        "institucionais não padronizados, e ao aplicar métodos rigorosos de "
        "detecção de quebras estruturais e validação temporal de modelos "
        "preditivos a um fenômeno recente e em rápida transformação regulatória."
    ), recuo_primeira_linha=1.25)

    subtitulo(doc, "1.5 Estrutura deste Relatório")
    paragrafo(doc, (
        "Este relatório está organizado em quatro capítulos. O Capítulo 2 "
        "apresenta a fundamentação teórica preliminar, reunindo evidências "
        "nacionais e internacionais sobre apostas online e vulnerabilidade "
        "financeira, o marco regulatório brasileiro e a literatura metodológica "
        "sobre detecção de quebras estruturais. O Capítulo 3 detalha a "
        "metodologia planejada e o estado atual de sua implementação técnica. O "
        "Capítulo 4 apresenta o cronograma oficial de entregas e o mapeamento das "
        "atividades técnicas a cada etapa. As Referências reúnem as obras citadas."
    ), recuo_primeira_linha=1.25)


# --- Capítulo 2 --------------------------------------------------------- #

def construir_fundamentacao(doc):
    titulo_capitulo(doc, "2 FUNDAMENTAÇÃO TEÓRICA PRELIMINAR")
    paragrafo(doc, (
        "Esta seção reúne, de forma preliminar, a literatura acadêmica e "
        "institucional identificada até o momento desta entrega. Trata-se de um "
        "levantamento inicial — a revisão de literatura será aprofundada até a "
        "Segunda Entrega, conforme indicado no Capítulo 4."
    ), recuo_primeira_linha=1.25)

    subtitulo(doc, "2.1 Apostas Online e Vulnerabilidade Financeira: Evidências Internacionais")
    paragrafo(doc, (
        "A literatura econômica internacional recente oferece evidência causal "
        "sobre o efeito da legalização de apostas esportivas online no "
        "endividamento das famílias. Baker et al. (2024), em estudo do National "
        "Bureau of Economic Research baseado em dados de transações bancárias e em "
        "desenho de diferenças em diferenças escalonado, estimam que a legalização "
        "das apostas esportivas online eleva o saldo devedor de cartão de crédito "
        "em torno de 11 a 12%, com o efeito concentrado em domicílios já "
        "financeiramente restritos: aumento do saldo devedor, redução do crédito "
        "disponível e maior frequência de uso de cheque especial. Os autores "
        "constatam ainda que apostadores frequentes destinam fração maior de sua "
        "renda às apostas e substituem poupança e investimentos de valor esperado "
        "positivo por apostas de risco, com efeitos desproporcionalmente maiores "
        "entre domicílios de menor escolaridade."
    ), recuo_primeira_linha=1.25)

    subtitulo(doc, "2.2 Apostas Online e Endividamento das Famílias Brasileiras: Evidências Nacionais")
    paragrafo(doc, (
        "No contexto brasileiro, três frentes de evidência convergem para o "
        "mesmo diagnóstico. Em primeiro lugar, o Banco Central do Brasil (2024) "
        "dimensionou pela primeira vez o mercado de apostas online no país e sua "
        "relação com o fluxo de pagamentos via Pix de beneficiários de programas "
        "sociais. Em segundo lugar, a Confederação Nacional do Comércio, por meio "
        "de sua Gerência Executiva de Análise e Desenvolvimento Econômico, aplicou "
        "metodologia de diferenças em diferenças sobre dados da Pesquisa de "
        "Endividamento e Inadimplência do Consumidor para estimar que as apostas "
        "online elevaram o número de famílias em inadimplência severa em "
        "aproximadamente 270 mil no período analisado (CNC, 2026) — abordagem "
        "quase-experimental que dialoga diretamente com a estratégia de detecção "
        "de quebras estruturais adotada neste trabalho. Em terceiro lugar, na "
        "literatura acadêmica nacional, Santos, Coelho e Bernardes (2025) analisam, "
        "por metodologia bibliográfica e documental, a relação entre apostas "
        "online e o superendividamento do consumidor brasileiro à luz do marco "
        "legal vigente, reforçando a pertinência de se investigar o fenômeno sob a "
        "ótica dos marcos regulatórios recentes."
    ), recuo_primeira_linha=1.25)

    subtitulo(doc, "2.3 O Marco Regulatório Brasileiro")
    paragrafo(doc, (
        "Três marcos regulatórios com data definida delimitam o período de "
        "interesse deste trabalho e serão utilizados, na etapa de detecção de "
        "quebras estruturais, exclusivamente como referência de validação externa "
        "— nunca como informação de entrada para os algoritmos de segmentação."
    ), recuo_primeira_linha=1.25)
    tabela_simples(
        doc,
        ["Marco", "Data", "Conteúdo"],
        [
            ["Regulamentação do mercado", "Janeiro de 2025",
             "Entrada em vigor da exigência de autorização das operadoras (Lei n.º 14.790/2023)."],
            ["Restrição a beneficiários", "Outubro de 2025",
             "Portaria SPA/MF n.º 2.217/2025 e Instrução Normativa SPA/MF n.º 22/2025 vedam a "
             "participação de beneficiários do Bolsa Família e do BPC."],
            ["Suspensão parcial judicial", "Dezembro de 2025",
             "Decisão do STF suspende as obrigações de bloqueio e encerramento de contas "
             "existentes, mantendo a vedação a novos cadastros."],
        ],
    )

    subtitulo(doc, "2.4 Detecção de Quebras Estruturais em Séries Temporais Econômicas")
    paragrafo(doc, (
        "A detecção de pontos de mudança (change points) em séries temporais sem "
        "conhecimento prévio de sua localização é um problema consolidado na "
        "literatura estatística. Truong, Oudre e Vayatis (2020) sistematizam os "
        "principais métodos offline de detecção de quebras estruturais, entre "
        "eles os algoritmos de programação dinâmica com penalização (Pelt) e de "
        "segmentação binária (Binseg), adotados na implementação técnica deste "
        "trabalho (Capítulo 3) por meio da biblioteca ruptures."
    ), recuo_primeira_linha=1.25)


# --- Capítulo 3 --------------------------------------------------------- #

def construir_metodologia(doc):
    titulo_capitulo(doc, "3 METODOLOGIA")

    subtitulo(doc, "3.1 Tipo de Pesquisa")
    paragrafo(doc, (
        "Trata-se de pesquisa quantitativa, de natureza descritiva, preditiva e "
        "documental. O trabalho não formula propostas de política pública, não "
        "avalia a eficácia da regulação vigente e não estabelece relações "
        "causais em nível individual — seu compromisso é levantar, integrar, "
        "tratar, modelar e projetar, a partir exclusivamente de fontes públicas "
        "de acesso programático, o que dispensa coleta primária."
    ), recuo_primeira_linha=1.25)

    subtitulo(doc, "3.2 Recorte e Unidade de Análise")
    paragrafo(doc, (
        "A unidade de análise principal é o Brasil: a modelagem estatística, a "
        "detecção de quebras estruturais e a projeção de cenários operam sobre "
        "séries temporais nacionais mensais. Duas camadas de aprofundamento "
        "exploratório e descritivo complementam esse núcleo."
    ), recuo_primeira_linha=1.25)
    tabela_simples(
        doc,
        ["Camada", "Unidade", "Tratamento analítico"],
        [
            ["Núcleo", "Brasil", "Análise exploratória completa, detecção de quebras, modelagem "
             "preditiva comparativa e projeção de cenários."],
            ["Recorte 1", "27 Unidades da Federação", "Indicador composto de exposição e ordenamento "
             "comparativo. Sem projeção de cenários."],
            ["Recorte 2", "Municípios da UF identificada", "Caracterização descritiva de um conjunto "
             "reduzido de municípios. Sem modelagem."],
        ],
    )
    paragrafo(doc, (
        "Não há divulgação regular de dados de apostas por Unidade da Federação "
        "ou por município: o indicador territorial é construído por aproximação, "
        "a partir de fontes indiretas, e apresentado com essa qualificação "
        "explícita em toda saída do pipeline — trata-se de limitação declarada, "
        "não de resultado a ser afirmado como medição oficial."
    ), recuo_primeira_linha=1.25)

    subtitulo(doc, "3.3 Fontes de Dados e Estratégia de Coleta")
    paragrafo(doc, (
        "Todas as fontes são públicas. Priorizou-se, de forma estrita, o acesso "
        "por interface de programação de aplicação (API) oficial; onde não há API "
        "estável confirmada, adotou-se a leitura de exportação manual oficial em "
        "vez de coleta automatizada da interface (scraping), decisão registrada "
        "em ata de decisão técnica (ADR) própria do projeto."
    ), recuo_primeira_linha=1.25)
    tabela_simples(
        doc,
        ["Fonte", "Camada", "Estratégia de acesso"],
        [
            ["SGS — Sistema Gerenciador de Séries Temporais (BCB)", "Núcleo", "API oficial (implementada e testada)"],
            ["IBGE — API SIDRA", "Núcleo", "API oficial (implementada e testada)"],
            ["Portal da Transparência / CadÚnico", "Recortes 1 e 2", "API oficial com token (implementada)"],
            ["PEIC — CNC", "Núcleo", "Exportação manual normalizada"],
            ["EPAE — BCB", "Núcleo", "Exportação manual normalizada"],
            ["Painéis semestrais SPA/MF", "Núcleo", "Exportação manual normalizada"],
            ["Google Trends", "Recortes 1 e 2", "Exportação CSV nativa da própria ferramenta"],
            ["ESTBAN — BCB", "Recorte 2", "Exportação manual normalizada"],
        ],
    )

    subtitulo(doc, "3.4 Tratamento e Análise Exploratória")
    paragrafo(doc, (
        "As séries brutas são deflacionadas pelo Índice Nacional de Preços ao "
        "Consumidor Amplo (IPCA), com ausências curtas (até dois meses "
        "consecutivos, no interior da série) imputadas por interpolação linear e "
        "ausências estruturais (início tardio de série) preservadas como tal, "
        "nunca extrapoladas retroativamente. Valores atípicos são sinalizados, "
        "não removidos automaticamente. A análise exploratória inclui testes de "
        "estacionariedade (Dickey-Fuller Aumentado e Kwiatkowski-Phillips-Schmidt-"
        "Shin), decomposição sazonal e correlação cruzada com defasagens de zero "
        "a doze meses entre o indicador de exposição a apostas e os indicadores "
        "de endividamento."
    ), recuo_primeira_linha=1.25)

    subtitulo(doc, "3.5 Detecção de Quebras Estruturais e Validação Externa")
    paragrafo(doc, (
        "Os algoritmos de segmentação são aplicados às séries financeiras "
        "nacionais sem qualquer informação prévia sobre as datas dos marcos "
        "regulatórios. As quebras detectadas são, em etapa subsequente e "
        "separada, confrontadas com o calendário regulatório apresentado na "
        "Seção 2.3 — constituindo o teste de validade externa da base."
    ), recuo_primeira_linha=1.25)

    subtitulo(doc, "3.6 Modelagem Preditiva e Validação")
    paragrafo(doc, (
        "A comparação de modelos segue uma escada que parte de referências "
        "ingênuas (naive e naive sazonal, obrigatórias como piso de comparação), "
        "passa por um modelo econométrico com variável exógena (SARIMAX) e "
        "alcança uma abordagem de aprendizado de máquina sobre variáveis de "
        "defasagem. A validação é feita exclusivamente por origem móvel "
        "(rolling origin / walk-forward): nenhum modelo é avaliado por divisão "
        "aleatória entre treino e teste, dado tratar-se de séries temporais. A "
        "comparação final entre modelos reporta simultaneamente uma métrica de "
        "erro relativo (Mean Absolute Scaled Error ou Symmetric Mean Absolute "
        "Percentage Error) e o resultado do teste estatístico de diferença de "
        "desempenho de Diebold e Mariano, com correção de pequena amostra de "
        "Harvey, Leybourne e Newbold."
    ), recuo_primeira_linha=1.25)

    subtitulo(doc, "3.7 Projeção de Cenários e Custo de Oportunidade Patrimonial")
    paragrafo(doc, (
        "O horizonte principal de projeção é de 24 meses, sob três cenários — "
        "inercial, de contenção e de expansão —, todos com intervalo de predição "
        "explícito. O cenário inercial é rotulado como condicional à manutenção "
        "do ambiente regulatório vigente, premissa que a própria cronologia "
        "recente contraria; horizontes superiores a 24 meses, quando produzidos, "
        "são rotulados como exercício de cenário, nunca como previsão. O custo de "
        "oportunidade patrimonial associado ao redirecionamento do gasto médio em "
        "apostas é estimado por simulação estocástica de Monte Carlo, com semente "
        "aleatória fixa para garantir reprodutibilidade exata do resultado."
    ), recuo_primeira_linha=1.25)

    subtitulo(doc, "3.8 Ferramentas e Tecnologias")
    paragrafo(doc, (
        "O pipeline é implementado em linguagem Python, com pandas para "
        "manipulação de dados, statsmodels para testes estatísticos e modelos "
        "econométricos, ruptures para detecção de quebras estruturais, "
        "scikit-learn para o modelo de aprendizado de máquina, matplotlib para "
        "geração de figuras e Streamlit para o painel interativo de consulta aos "
        "resultados. Um servidor de Model Context Protocol (MCP) local padroniza, "
        "por meio de seis módulos de conhecimento (skills), como cada etapa do "
        "trabalho deve ser conduzida, impedindo que etapas de modelagem, "
        "avaliação ou redação avancem sem consultar previamente o contexto e as "
        "restrições da proposta."
    ), recuo_primeira_linha=1.25)

    subtitulo(doc, "3.9 Estado Atual da Implementação")
    paragrafo(doc, (
        "Até o momento desta Primeira Entrega, o grupo concluiu a especificação "
        "completa do projeto por desenvolvimento orientado por especificações, a "
        "definição da arquitetura técnica e a implementação e validação "
        "automatizada de 27 módulos de software, organizados nas cinco frentes "
        "descritas nas seções anteriores, com 161 testes automatizados passando — "
        "159 sem dependência de rede e dois validados contra as interfaces reais "
        "do Banco Central e do IBGE. Três das oito fontes de dados (Banco Central "
        "via SGS, IBGE via SIDRA e Portal da Transparência) já estão integradas a "
        "interfaces de programação de aplicação oficiais reais, com os "
        "identificadores de série tecnicamente verificados, não presumidos. A "
        "execução do pipeline completo sobre dados reais coletados, incluindo as "
        "cinco fontes de exportação manual, está planejada para o período que "
        "antecede a Segunda Entrega, conforme detalhado no Capítulo 4."
    ), recuo_primeira_linha=1.25)

    subtitulo(doc, "3.10 Princípios Metodológicos Inegociáveis")
    paragrafo(doc, "O trabalho é conduzido sob os seguintes princípios, formalmente registrados e "
                    "não sujeitos a exceção sem revisão documentada:", recuo_primeira_linha=1.25)
    for texto in [
        "Reprodutibilidade obrigatória: sementes aleatórias fixas em todo processo estocástico, ambiente declarado com versões de biblioteca fixadas;",
        "Ausência de vazamento de dados entre treino e teste, com validação exclusivamente por origem móvel;",
        "Data de corte fixa da base nacional em 31 de dezembro de 2025, cobrindo integralmente os três marcos regulatórios;",
        "Preferência estrita por fontes com interface de programação de aplicação oficial, com coleta automatizada da interface (scraping) tratada como último recurso, sempre documentado;",
        "Fidelidade ao escopo por camada: projeção de cenários restrita ao núcleo nacional; o Recorte 1 nunca é projetado e o Recorte 2 nunca é modelado.",
    ]:
        paragrafo(doc, f"• {texto}", recuo_esquerdo=1.25)


# --- Capítulo 4 --------------------------------------------------------- #

def construir_cronograma(doc):
    titulo_capitulo(doc, "4 CRONOGRAMA")
    paragrafo(doc, (
        "O cronograma institucional organiza o Trabalho de Conclusão de Curso em "
        "três entregas quinzenais, apresentadas no Quadro 1."
    ), recuo_primeira_linha=1.25)
    tabela_simples(
        doc,
        ["Entrega", "Início da quinzena", "Vencimento", "Carência"],
        [
            ["Primeira Entrega", "24/08/2026 (Quinzena 2)", "01/09/2026, 23:59", "06/09/2026, 23:59"],
            ["Segunda Entrega", "05/10/2026 (Quinzena 5)", "13/10/2026, 23:59", "18/10/2026, 23:59"],
            ["Entrega Final", "02/11/2026 (Quinzena 7)", "10/11/2026, 23:59", "15/11/2026, 23:59"],
        ],
    )
    paragrafo(doc, "Quadro 1 — Cronograma institucional de entregas.", tamanho=10,
              espacamento_simples=True, alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph()

    paragrafo(doc, (
        "O Quadro 2 mapeia as principais atividades técnicas planejadas para cada "
        "entrega, com base no plano de execução do projeto."
    ), recuo_primeira_linha=1.25)
    tabela_simples(
        doc,
        ["Entrega", "Atividades técnicas planejadas"],
        [
            ["Primeira Entrega (concluída nesta versão)",
             "Especificação do projeto; arquitetura técnica; implementação e testes dos módulos de "
             "ingestão, tratamento, modelagem, avaliação e relatório com dados sintéticos; integração "
             "real com três das oito fontes de dados; fundamentação teórica preliminar."],
            ["Segunda Entrega",
             "Coleta das cinco fontes restantes (exportação manual); execução do pipeline completo "
             "sobre dados reais; primeiros resultados de detecção de quebras estruturais e de "
             "comparação de modelos; aprofundamento da revisão de literatura."],
            ["Entrega Final",
             "Projeção de cenários e simulação de custo patrimonial sobre dados reais; indicador "
             "territorial consolidado; redação completa da monografia; publicação da base de dados, "
             "do código e do painel interativo sob princípios de dados abertos."],
        ],
    )
    paragrafo(doc, "Quadro 2 — Mapeamento de atividades técnicas por entrega.", tamanho=10,
              espacamento_simples=True, alinhamento=WD_ALIGN_PARAGRAPH.LEFT)


# --- Referências ---------------------------------------------------------- #

def construir_referencias(doc):
    titulo_capitulo(doc, "REFERÊNCIAS")

    refs = [
        "BAKER, Scott R.; BALTHROP, Justin; JOHNSON, Mark J.; KOTTER, Jason D.; "
        "PISCIOTTA, Kevin. Gambling Away Stability: Sports Betting's Impact on "
        "Vulnerable Households. NBER Working Paper Series, n. 33108, Cambridge, MA: "
        "National Bureau of Economic Research, 2024.",

        "BANCO CENTRAL DO BRASIL. Estudo Especial n. 119: mercado de apostas "
        "online e perfil dos apostadores. Brasília: BCB, ago. 2024.",

        "BRASIL. Lei n.º 14.790, de 29 de dezembro de 2023. Dispõe sobre a "
        "regulamentação da modalidade lotérica de apostas de quota fixa. Diário "
        "Oficial da União, Brasília, DF.",

        "BRASIL. Secretaria de Prêmios e Apostas. Portaria SPA/MF n.º 2.217, de "
        "2025. Diário Oficial da União, Brasília, DF.",

        "BRASIL. Secretaria de Prêmios e Apostas. Instrução Normativa SPA/MF n.º "
        "22, de 2025. Diário Oficial da União, Brasília, DF.",

        "COMITÊ NACIONAL DE SECRETÁRIOS DE FAZENDA (COMSEFAZ); CENTRO "
        "INTERNACIONAL CELSO FURTADO. Boletim Fiscal dos Estados Brasileiros. "
        "3. ed. [S. l.]: Comsefaz, ago. 2026.",

        "CONFEDERAÇÃO NACIONAL DO COMÉRCIO DE BENS, SERVIÇOS E TURISMO (CNC). "
        "Gerência Executiva de Análise e Desenvolvimento Econômico (Geade). "
        "Impactos das apostas online (bets) no endividamento e na inadimplência "
        "das famílias brasileiras. Brasília: CNC, 2026.",

        "SANTOS, Gabrielly Cordeiro dos; COELHO, Ivana Lara Ribeiro; BERNARDES, "
        "Rochele Juliane Lima Firmeza. Entre a diversão e a ruína: a influência "
        "das apostas online/BETS no endividamento excessivo do brasileiro. "
        "Revista Ibero-Americana de Humanidades, Ciências e Educação, v. 11, "
        "n. 4, p. 3396-3420, abr. 2025.",

        "SECRETARIA DE PRÊMIOS E APOSTAS (SPA/MF). Boletim semestral do mercado "
        "de apostas — 1º semestre de 2025. Brasília: Ministério da Fazenda, 2025.",

        "TRUONG, Charles; OUDRE, Laurent; VAYATIS, Nicolas. Selective review of "
        "offline change point detection methods. Signal Processing, v. 167, "
        "2020.",

        "UNIVERSIDADE VIRTUAL DO ESTADO DE SÃO PAULO. Manual de Normalização de "
        "Trabalhos Acadêmicos da UNIVESP. São Paulo: UNIVESP, [s.d.].",

        "ASSOCIAÇÃO BRASILEIRA DE NORMAS TÉCNICAS (ABNT). NBR 14724: informação "
        "e documentação — trabalhos acadêmicos — apresentação. Rio de Janeiro: "
        "ABNT, 2011.",

        "ASSOCIAÇÃO BRASILEIRA DE NORMAS TÉCNICAS (ABNT). NBR 6023: informação "
        "e documentação — referências — elaboração. Rio de Janeiro: ABNT, 2018.",
    ]
    refs_ordenadas = sorted(refs, key=lambda r: r.upper())
    for ref in refs_ordenadas:
        referencia(doc, ref)


def main():
    doc = Document()
    configurar_documento(doc)

    construir_capa(doc)
    construir_folha_de_rosto(doc)
    construir_ficha(doc)
    construir_resumo(doc)
    construir_sumario(doc)
    construir_introducao(doc)
    construir_fundamentacao(doc)
    construir_metodologia(doc)
    construir_cronograma(doc)
    construir_referencias(doc)

    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    doc.save(SAIDA)
    print(f"Documento gerado em: {SAIDA}")


if __name__ == "__main__":
    main()
